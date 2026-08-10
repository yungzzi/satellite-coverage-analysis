# -*- coding: utf-8 -*-
"""
validate_selection.py —— 选星模块回归自检

Q104 教训：AppTest 只能渲染默认第一个页签、点不到⑤页按钮，之前仿真脚本又自己
手写了 windows_by_sat，导致"测试通过但界面 NameError"。本自检统一走
cov.analyze_collection_pipeline（⑤页按钮调用的同一个函数），堵住这个漏洞。

运行：.venv/Scripts/python src\validate_selection.py
"""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import coverage_analysis as cov
import satellite_loader as sl
import satellite_selection as ss

PASS = 0
FAIL = 0


def check(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [PASS] {name}{('  ' + detail) if detail else ''}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name}{('  ' + detail) if detail else ''}")


def jilin_recs():
    cols = sl.load_builtin_collections()
    jl = next(c for c in cols if "吉林一号全量" in c["name"])
    return sl.resolve_collection(jl)


def make_cfg(region_name="杭州5km"):
    return {
        "region_name": region_name,
        "region_mode": "circle",
        "tle_file": "自检（吉林一号测试快照）",
        "fov_half_angle_deg": 30.0,
        "min_elevation_deg": 10.0,
        "step_minutes": 1,
        "refine_step_seconds": 10,
        "grid_spacing_deg": 0.2,
        "target": {"lat": 30.25, "lon": 120.15, "radius_km": 5.0},
        "save_gantt": False,
        "save_plots": False,
        "max_gantt_rows": 15,
    }


def main():
    print("== 1. 覆盖分析共用入口（⑤页按钮同款路径，防 Q104 回归）==")
    recs = jilin_recs()
    cfg = make_cfg()
    region = {"mode": "circle", "lat": 30.25, "lon": 120.15, "radius_km": 5.0}
    result = cov.analyze_collection_pipeline(
        recs, region, cfg, "2026-08-03 00:00:00", "2026-08-04 00:00:00")
    check("入口返回结果", result is not None)
    check("windows_by_sat 与卫星数一致",
          result is not None and len(result["windows_by_sat"]) == result["n_sats"] == len(recs),
          f"{len(result['windows_by_sat']) if result else 0} vs {len(recs)}")
    check("窗口数与明细一致",
          result is not None and result["n_win"] == sum(
              len(w) for w in result["windows_by_sat"]),
          f"{result['n_win'] if result else 0}")
    check("四件套输出齐全",
          result is not None and len(result["dfs"]) == 4)

    print("== 2. 选星：贪心 vs 整数规划（可满足需求，MRT=180）==")
    if result:
        sel = ss.run_selection(result["sats"], result["windows_by_sat"],
                               result["t0"], result["t1"],
                               step_minutes=1, k_min=1, mrt_minutes=180,
                               method="both", time_limit=120, req_mode="revisit")
        g = sel["greedy"]; i = sel["ilp"]
        check("贪心满足需求", g["eval"]["satisfied"],
              f"MRT={g['eval']['mrt_actual_minutes']:.0f}分")
        check("ILP 证明最优", i is not None and i["status"] == "Optimal",
              i["status"] if i else "无")
        check("两算法数量一致", len(g["selected"]) == len(i["selected"]),
              f"贪心 {len(g['selected'])} vs ILP {len(i['selected'])}")
        check("采用 ILP 结果", sel["algo"] == "ilp" and sel["eval"]["satisfied"])

        print("== 3. 需求不可满足时如实报告（MRT=30）==")
        hard = ss.run_selection(result["sats"], result["windows_by_sat"],
                                result["t0"], result["t1"],
                                step_minutes=1, k_min=1, mrt_minutes=30,
                                method="both", time_limit=60, req_mode="revisit")
        check("不硬凑：如实报未满足", not hard["eval"]["satisfied"],
              f"MRT={hard['eval']['mrt_actual_minutes']:.0f}分")
        check("空窗表非空", len(hard["gaps"]) > 0)

        print("== 4. 选星 Excel 导出 ==")
        out = os.path.join(tempfile.gettempdir(), "sel_validate.xlsx")
        ss.export_selection_xlsx(sel, out, region_name="杭州5km")
        check("Excel 非空", os.path.getsize(out) > 0, f"{os.path.getsize(out)} B")
        import openpyxl
        _wb = openpyxl.load_workbook(out, read_only=True)
        _sheets = _wb.sheetnames
        check("Excel 含全部最优方案表", "全部最优方案" in _sheets, str(_sheets))
        check("Excel 说明表含需求模式", any(
            "需求模式" in str(c.value) for row in _wb["说明"].iter_rows() for c in row if c.value))
        _wb.close()

    print("== 5. 需求双模式 + 目标模式② + 全部最优方案 ==")
    if result:
        # 连续覆盖（吉林一号纯低轨 → 物理不可行，如实报告上限）
        cont = ss.run_selection(result["sats"], result["windows_by_sat"],
                                result["t0"], result["t1"],
                                step_minutes=1, k_min=1, mrt_minutes=30,
                                method="greedy", req_mode="continuous")
        check("连续模式：物理不可行如实报告",
              cont["req_mode"] == "continuous" and not cont["eval"]["satisfied"]
              and cont.get("feasibility") is not None
              and not cont["feasibility"]["possible"],
              f"possible={cont.get('feasibility', {}).get('possible')}")
        check("连续模式：给出物理上限数字",
              cont.get("feasibility", {}).get("full_max_gap_minutes", 0) > 0,
              str(cont.get("feasibility", {}).get("full_max_gap_minutes")))
        # 重访模式兼容（原有语义）
        rev = ss.run_selection(result["sats"], result["windows_by_sat"],
                               result["t0"], result["t1"],
                               step_minutes=1, k_min=1, mrt_minutes=180,
                               method="greedy", req_mode="revisit")
        check("重访模式：字段与判定正常",
              rev["req_mode"] == "revisit" and rev["effective_mrt_minutes"] == 180.0
              and rev["eval"]["satisfied"])
        # 目标模式②：给定 N 颗最大化覆盖率
        mc = ss.run_selection(result["sats"], result["windows_by_sat"],
                              result["t0"], result["t1"],
                              step_minutes=1, k_min=1, mrt_minutes=30,
                              method="greedy", req_mode="continuous",
                              objective="max_coverage", n_sats_limit=5)
        check("目标模式②：选中数≤N 且覆盖率合理",
              mc["objective"] == "max_coverage"
              and 0 < len(mc["selected"]) <= 5
              and 0 < mc["eval"]["coverage_pct"] <= cont["full"]["coverage_pct"] + 1e-9,
              f"选 {len(mc['selected'])} 颗，覆盖率 {mc['eval']['coverage_pct']:.2f}%")
        # 全部最优方案（合成 4 颗全程星）
        import numpy as np
        _m = 10
        _sc = [np.ones(_m, dtype=bool) for _ in range(4)]
        _ss = [type("S", (), {"name": f"FULL-{j}"})() for j in range(4)]
        _s1, _n1 = ss.enumerate_optimal_solutions(_sc, _ss, 1, 0, 1)
        _s2, _n2 = ss.enumerate_optimal_solutions(_sc, _ss, 2, 0, 1)
        check("全部最优方案 K=1 → 4 种", len(_s1) == 4, str(len(_s1)))
        check("全部最优方案 K=2 → 6 种", len(_s2) == 6, str(len(_s2)))

    print("== 6. 预算型任务（MMRT/MART）+ 五任务字段 + 尽力而为说明 ==")
    if result:
        from datetime import datetime, timedelta, timezone
        import numpy as np
        _m2 = 12
        _sc2 = [
            np.array([1,1,1,1,1,1,0,0,0,0,0,0], bool),
            np.array([0,0,0,0,1,1,1,1,1,1,0,0], bool),
            np.array([0,0,0,0,0,0,0,1,1,1,1,1], bool),
            np.array([1,1,1,0,0,0,0,0,0,0,0,0], bool),
            np.array([0,0,0,0,0,0,1,1,1,1,1,1], bool),
            np.array([0,0,0,0,0,0,0,0,0,0,0,1], bool),
        ]
        _ss2 = [type("S", (), {"name": f"T{j}"})() for j in range(len(_sc2))]
        _t0 = datetime(2026, 1, 1, tzinfo=timezone.utc)
        _t1 = _t0 + timedelta(minutes=_m2)
        s_mm, _c1 = ss.select_min_gap(_sc2, 2, 1, 1, _t0, _t1, mode="max")
        s_ma, _c2 = ss.select_min_gap(_sc2, 2, 1, 1, _t0, _t1, mode="avg")
        check("MMRT 贪心能选出星", 0 < len(s_mm) <= 2, f"选 {len(s_mm)} 颗")
        check("MART 贪心能选出星", 0 < len(s_ma) <= 2, f"选 {len(s_ma)} 颗")
        wbs2 = []
        for c in _sc2:
            wins = []
            i2 = 0
            while i2 < len(c):
                if c[i2]:
                    j2 = i2
                    while j2 < len(c) and c[j2]:
                        j2 += 1
                    wins.append((_t0 + timedelta(minutes=i2),
                                 _t0 + timedelta(minutes=j2), 1.0))
                    i2 = j2
                else:
                    i2 += 1
            wbs2.append(wins)
        for obj, req, tag in [("max_coverage", "any", "③"),
                              ("min_max_gap", "any", "④"),
                              ("min_avg_gap", "any", "⑤")]:
            r5 = ss.run_selection(_ss2, wbs2, _t0, _t1, step_minutes=1, k_min=1,
                                  mrt_minutes=180, method="greedy",
                                  req_mode=req, objective=obj, n_sats_limit=2)
            check(f"任务{tag}：字段与标签正确",
                  r5["req_mode"] == "any" and r5["objective"] == obj
                  and r5["algo"] == "greedy-budget"
                  and "目标模式" in r5["summary"],
                  f"obj={r5['objective']} algo={r5['algo']}")
        cont2 = ss.run_selection(result["sats"], result["windows_by_sat"],
                                 result["t0"], result["t1"],
                                 step_minutes=1, k_min=1, mrt_minutes=30,
                                 method="greedy", req_mode="continuous")
        check("未满足时摘要含尽力而为说明",
              "[尽力而为]" in cont2["summary"])

    print("== 7. PSCLP：覆盖率阈值 + 最少颗数 ==")
    if result:
        from datetime import datetime, timedelta, timezone
        import numpy as np
        _m3 = 12
        _sc3 = [
            np.array([1,1,1,1,1,1,0,0,0,0,0,0], bool),
            np.array([0,0,0,0,1,1,1,1,1,1,0,0], bool),
            np.array([0,0,0,0,0,0,0,1,1,1,1,1], bool),
            np.array([1,1,1,0,0,0,0,0,0,0,0,0], bool),
            np.array([0,0,0,0,0,0,1,1,1,1,1,1], bool),
            np.array([0,0,0,0,0,0,0,0,0,0,0,1], bool),
        ]
        _ss3 = [type("S", (), {"name": f"P{j}"})() for j in range(len(_sc3))]
        _t0 = datetime(2026, 1, 1, tzinfo=timezone.utc)
        _t1 = _t0 + timedelta(minutes=_m3)
        s50, c50 = ss.greedy_select_coverage(_sc3, 1, 50.0, 1)
        s100, c100 = ss.greedy_select_coverage(_sc3, 1, 100.0, 1)
        check("PSCLP P=50% 用星少于等于 P=100%", len(s50) <= len(s100),
              f"{len(s50)} vs {len(s100)}")
        check("PSCLP P=50% 覆盖率达标", (c50 >= 1).mean() * 100 >= 50.0 - 1e-9,
              f"{(c50 >= 1).mean() * 100:.1f}%")
        check("PSCLP P=100% 覆盖率达标", (c100 >= 1).mean() * 100 >= 99.9,
              f"{(c100 >= 1).mean() * 100:.1f}%")
        wbs3 = []
        for c in _sc3:
            wins = []
            i2 = 0
            while i2 < len(c):
                if c[i2]:
                    j2 = i2
                    while j2 < len(c) and c[j2]:
                        j2 += 1
                    wins.append((_t0 + timedelta(minutes=i2),
                                 _t0 + timedelta(minutes=j2), 1.0))
                    i2 = j2
                else:
                    i2 += 1
            wbs3.append(wins)
        rp = ss.run_selection(_ss3, wbs3, _t0, _t1, step_minutes=1, k_min=1,
                              mrt_minutes=180, method="greedy",
                              req_mode="coverage_pct", coverage_pct=50.0)
        check("PSCLP run_selection 字段正确",
              rp["req_mode"] == "coverage_pct" and rp["coverage_threshold_pct"] == 50.0
              and rp["eval"]["satisfied"])
        rp2 = ss.run_selection(result["sats"], result["windows_by_sat"],
                               result["t0"], result["t1"],
                               step_minutes=1, k_min=1, mrt_minutes=180,
                               method="greedy", req_mode="coverage_pct", coverage_pct=100.0)
        check("PSCLP P=100% 纯低轨如实不满足",
              not rp2["eval"]["satisfied"]
              and "PSCLP" in rp2["summary"])

    print("== 8. 新算法回归：PSCLP-ILP / 任务③④ ILP / K>=2 贪心修复 ==")
    try:
        cols8 = sl.load_builtin_collections()
        col8 = next(c for c in cols8 if "中国遥感全量" in c["name"])
        recs8 = sl.resolve_collection(col8)
        cfg8 = make_cfg("杭州5km")
        region8 = {"mode": "circle", "lat": 30.25, "lon": 120.15, "radius_km": 5.0}
        res8 = cov.analyze_collection_pipeline(recs8, region8, cfg8,
                                               "2026-08-03 00:00:00",
                                               "2026-08-03 08:00:00")
        r_a = ss.run_selection(res8["sats"], res8["windows_by_sat"],
                               res8["t0"], res8["t1"],
                               step_minutes=1, k_min=2, mrt_minutes=1.0,
                               method="both", req_mode="coverage_pct", coverage_pct=100.0)
        check("PSCLP K=2 P=100 贪心修复（≥2 颗且满足）",
              r_a["greedy"] is not None and len(r_a["greedy"]["selected"]) >= 2
              and r_a["greedy"]["eval"]["satisfied"],
              f"贪心 {len(r_a['greedy']['selected']) if r_a['greedy'] else 0} 颗")
        check("PSCLP K=2 P=100 ILP 精确最优",
              r_a["ilp"] is not None and r_a["ilp"]["status"] == "Optimal"
              and len(r_a["ilp"]["selected"]) == 2 and r_a["ilp"]["eval"]["satisfied"],
              f"ILP {len(r_a['ilp']['selected']) if r_a['ilp'] else 0} 颗")
        check("PSCLP K=2 P=100 采用 ILP 结果",
              r_a["algo"] == "ilp" and r_a["eval"]["satisfied"])
        check("PSCLP K=2 P=100 最优方案枚举",
              len(r_a.get("optimal_solutions") or []) >= 2,
              f"{len(r_a.get('optimal_solutions') or [])} 种")
        r_c = ss.run_selection(res8["sats"], res8["windows_by_sat"],
                               res8["t0"], res8["t1"],
                               step_minutes=1, k_min=1, mrt_minutes=1.0,
                               method="both", req_mode="any", objective="max_coverage",
                               n_sats_limit=5)
        check("任务③ ILP 精确最优（覆盖率100%）",
              r_c["ilp"] is not None and r_c["ilp"]["status"] == "Optimal"
              and r_c["eval"]["coverage_pct"] >= 99.9,
              f"algo={r_c['algo']} cov={r_c['eval']['coverage_pct']:.1f}%")
        r_d = ss.run_selection(res8["sats"], res8["windows_by_sat"],
                               res8["t0"], res8["t1"],
                               step_minutes=1, k_min=1, mrt_minutes=1.0,
                               method="both", req_mode="any", objective="min_max_gap",
                               n_sats_limit=3)
        check("任务④ ILP 精确最优（最长空窗0）",
              r_d["ilp"] is not None and r_d["ilp"]["status"] == "Optimal"
              and r_d["eval"]["mrt_actual_minutes"] <= 1.0,
              f"algo={r_d['algo']} 空窗={r_d['eval']['mrt_actual_minutes']:.0f}分")
        r_e = ss.run_selection(res8["sats"], res8["windows_by_sat"],
                               res8["t0"], res8["t1"],
                               step_minutes=1, k_min=1, mrt_minutes=1.0,
                               method="both", req_mode="any", objective="min_avg_gap",
                               n_sats_limit=3)
        check("任务⑤ 保持贪心并注明分式目标",
              r_e["algo"] == "greedy-budget" and "分式目标" in r_e["algo_note"],
              r_e["algo_note"])
        r_k6 = ss.run_selection(res8["sats"], res8["windows_by_sat"],
                                res8["t0"], res8["t1"],
                                step_minutes=1, k_min=6, mrt_minutes=1.0,
                                method="both", req_mode="coverage_pct", coverage_pct=100.0)
        check("K=6 P=100 ILP 证明不可行 + 提示区分",
              not r_k6["eval"]["satisfied"]
              and r_k6.get("ilp") is not None and r_k6["ilp"]["status"] == "Infeasible"
              and "已证明" in r_k6.get("algo_note", ""),
              r_k6.get("algo_note", ""))
        check("两者对照：贪心对照行存在",
              r_a["algo"] == "ilp" and len(r_a["compare_rows"]) >= 3
              and any("贪心" in str(x.get("口径", "")) for x in r_a["compare_rows"]),
              f"{len(r_a['compare_rows'])} 行")
        check("精确评估：最终指标为窗口级（exact=True，空窗含秒）",
              r_a["eval"].get("exact") is True
              and all("duration_s" in g for g in r_a["gaps"]),
              f"gaps={len(r_a['gaps'])}")
        r_d2 = ss.run_selection(res8["sats"], res8["windows_by_sat"],
                                res8["t0"], res8["t1"],
                                step_minutes=2, k_min=1, mrt_minutes=2.0,
                                method="both", req_mode="any", objective="min_max_gap",
                                n_sats_limit=3)
        check("P2 修复：step>1 的 min_max_gap 正常出解",
              r_d2["eval"]["mrt_actual_minutes"] >= 0
              and r_d2["eval"].get("exact") is True,
              f"algo={r_d2['algo']} 空窗={r_d2['eval']['mrt_actual_minutes']:.1f}分")
    except Exception as _ex:
        check("新算法回归（中国遥感全量）", False, str(_ex))

    print("== 9. 健壮性：circle_overlap_ratio 纯 float 输入 ==")
    _r = cov.circle_overlap_ratio(50.0, 30.0, 20.0)
    check("float 输入部分相交不崩", isinstance(_r, float) and 0.0 <= _r <= 1.0,
          f"r={_r}")

    print("========================================")
    print(f"结果：{PASS} 项通过，{FAIL} 项失败")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
